
from pathlib import Path
import csv
import json

import pandas as pd
from openpyxl import load_workbook

# --------------------------------------------------
# 1. Project paths
# --------------------------------------------------

ROOT = Path(__file__).resolve().parent.parent

RAW_FILE = ROOT / "Data" / "Raw" / "online_retail_II.xlsx"
PROCESSED_DIR = ROOT / "Data" / "Processed"

PROCESSED_DIR.mkdir(parents=True, exist_ok=True)

if not RAW_FILE.exists():
    raise FileNotFoundError(
        f"Raw dataset not found: {RAW_FILE}\n"
        "Place online_retail_II.xlsx inside Data/Raw/"
    )

# --------------------------------------------------
# 2. Convert Excel worksheets to CSV using streaming
# --------------------------------------------------

workbook = load_workbook(
    RAW_FILE,
    read_only=True,
    data_only=True
)

expected_columns = [
    "invoice",
    "stockcode",
    "description",
    "quantity",
    "invoicedate",
    "unit_price",
    "customer_id",
    "country",
]

sheet_outputs = []

for sheet_name in workbook.sheetnames:
    sheet = workbook[sheet_name]

    output_name = (
        "online_retail_2009_2010.csv"
        if "2009" in sheet_name
        else "online_retail_2010_2011.csv"
    )

    output_path = PROCESSED_DIR / output_name

    rows_written = 0

    with output_path.open(
        "w", newline="", encoding="utf-8-sig"
    ) as file:
        writer = csv.writer(file)

        headers = next(
            sheet.iter_rows(min_row=1, max_row=1, values_only=True)
        )

        normalized_headers = [
            str(column).strip().lower().replace(" ", "_")
            for column in headers
        ]

        normalized_headers = [
            "unit_price" if col == "price" else col
            for col in normalized_headers
        ]

        if normalized_headers != expected_columns:
            raise ValueError(
                f"Unexpected columns in {sheet_name}: "
                f"{normalized_headers}"
            )

        writer.writerow(normalized_headers)

        for row in sheet.iter_rows(
            min_row=2, values_only=True
        ):
            writer.writerow(row)
            rows_written += 1

    sheet_outputs.append(output_name)

    print(
        f"Created {output_name}: "
        f"{rows_written:,} data rows"
    )

workbook.close()

# --------------------------------------------------
# 3. Process CSV files in chunks
# --------------------------------------------------

sales_path = PROCESSED_DIR / "online_retail_sales_initial.csv"

audit = {
    "source_file": RAW_FILE.name,
    "sheets_processed": sheet_outputs,
    "rows_read": 0,
    "rows_missing_customer_id": 0,
    "rows_missing_description": 0,
    "exact_duplicate_rows": 0,
    "cancellation_invoice_rows": 0,
    "negative_quantity_rows": 0,
    "zero_price_rows": 0,
    "negative_price_rows": 0,
    "initial_sales_rows": 0,
    "initial_sales_value": 0.0,
}

with sales_path.open(
    "w", newline="", encoding="utf-8-sig"
) as output_file:

    sales_writer = None

    for filename in sheet_outputs:
        csv_path = PROCESSED_DIR / filename

        for chunk in pd.read_csv(
            csv_path,
            chunksize=100_000,
            dtype={
                "invoice": "string",
                "stockcode": "string",
                "description": "string",
                "country": "string",
            },
        ):
            # Parse numeric columns safely
            chunk["quantity"] = pd.to_numeric(
                chunk["quantity"], errors="coerce"
            )

            chunk["unit_price"] = pd.to_numeric(
                chunk["unit_price"], errors="coerce"
            )

            chunk["customer_id"] = pd.to_numeric(
                chunk["customer_id"], errors="coerce"
            )

            chunk["invoicedate"] = pd.to_datetime(
                chunk["invoicedate"], errors="coerce"
            )

            audit["rows_read"] += len(chunk)
            audit["rows_missing_customer_id"] += int(
                chunk["customer_id"].isna().sum()
            )
            audit["rows_missing_description"] += int(
                chunk["description"].isna().sum()
            )
            audit["exact_duplicate_rows"] += int(
                chunk.duplicated().sum()
            )

            invoice_text = chunk["invoice"].fillna("")

            audit["cancellation_invoice_rows"] += int(
                invoice_text.str.startswith("C").sum()
            )
            audit["negative_quantity_rows"] += int(
                (chunk["quantity"] < 0).sum()
            )
            audit["zero_price_rows"] += int(
                (chunk["unit_price"] == 0).sum()
            )
            audit["negative_price_rows"] += int(
                (chunk["unit_price"] < 0).sum()
            )

            # Initial descriptive sales filter only.
            # We do not silently deduplicate the data.
            mask = (
                (chunk["quantity"] > 0)
                & (chunk["unit_price"] > 0)
                & (~invoice_text.str.startswith("C"))
            )

            sales = chunk.loc[mask].copy()

            sales["line_sales_value"] = (
                sales["quantity"] * sales["unit_price"]
            )

            if not sales.empty:
                if sales_writer is None:
                    sales.to_csv(
                        output_file,
                        index=False,
                        header=True,
                    )
                    sales_writer = True
                else:
                    sales.to_csv(
                        output_file,
                        index=False,
                        header=False,
                    )

                audit["initial_sales_rows"] += len(sales)
                audit["initial_sales_value"] += float(
                    sales["line_sales_value"].sum()
                )

# --------------------------------------------------
# 4. Save processing audit
# --------------------------------------------------

audit_path = PROCESSED_DIR / "processing_audit.json"

with audit_path.open("w", encoding="utf-8") as file:
    json.dump(audit, file, indent=4, default=str)

print("\nProcessing completed.")
print("Processed folder:", PROCESSED_DIR)
print("Initial sales file:", sales_path)
print("Audit report:", audit_path)
print(json.dumps(audit, indent=2))
